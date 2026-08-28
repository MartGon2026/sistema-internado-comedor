import secrets
import string
import unicodedata
import calendar
from datetime import date, datetime, timedelta
from django.db.models import Count, Q
from django.core.paginator import Paginator
from datetime import date, datetime
from django.shortcuts import render, redirect
from django.http import JsonResponse, HttpResponse
from django.template import loader
from django.db import IntegrityError
from django.db.models import Count, Q
from django.utils import timezone
from datetime import date, datetime, time
from datetime import datetime, timedelta

from django.db.models import Q
from django.utils import timezone

from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.utils import timezone

from openpyxl import Workbook

from datetime import date, datetime
from openpyxl import load_workbook, Workbook

from .models import Usuarios, Estudiantes, Asistencias, HistorialInternado, GuardaSeguridad
from .servicios.supabase_storage import subir_foto_estudiante, eliminar_foto_estudiante, subir_foto_guarda, eliminar_foto_guarda




def dashboard(request):

    hoy = timezone.localdate()

    # =====================================================
    # FILTRO POR FECHA
    # =====================================================

    desde = request.GET.get("desde") or hoy.isoformat()
    hasta = request.GET.get("hasta") or hoy.isoformat()

    try:
        fecha_desde = date.fromisoformat(desde)
        fecha_hasta = date.fromisoformat(hasta)

    except (ValueError, TypeError):
        fecha_desde = hoy
        fecha_hasta = hoy

    # Corregir las fechas si fueron seleccionadas al revés
    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde


    # =====================================================
    # DATOS GENERALES DE ESTUDIANTES
    # =====================================================

    total_estudiantes = Estudiantes.objects.count()

    estudiantes_dentro = Estudiantes.objects.filter(
        estado_internado__iexact="DENTRO"
    ).count()

    estudiantes_fuera = Estudiantes.objects.filter(
        estado_internado__iexact="FUERA"
    ).count()


    # =====================================================
    # ASISTENCIAS DEL COMEDOR EN EL RANGO
    # =====================================================

    asistencias = Asistencias.objects.filter(
        fecha__range=(fecha_desde, fecha_hasta),
        estado=True
    )

    asistencias_rango = asistencias.count()

    desayunos_servidos = asistencias.filter(
        tipo_comida__iexact="DESAYUNO"
    ).count()

    almuerzos_servidos = asistencias.filter(
        tipo_comida__iexact="ALMUERZO"
    ).count()

    cenas_servidas = asistencias.filter(
        tipo_comida__iexact="CENA"
    ).count()


    # =====================================================
    # MOVIMIENTOS DEL INTERNADO EN EL RANGO
    # =====================================================

    salidas_rango = HistorialInternado.objects.filter(
        fecha_salida__range=(fecha_desde, fecha_hasta)
    ).count()

    entradas_rango = HistorialInternado.objects.filter(
        fecha_entrada__range=(fecha_desde, fecha_hasta)
    ).count()


    # =====================================================
    # ESTUDIANTES QUE TODAVÍA NO HAN REGRESADO
    #
    # No se filtran por fecha porque interesa conocer
    # quiénes están fuera actualmente.
    # =====================================================

    estudiantes_faltantes = (
        HistorialInternado.objects
        .select_related("carnet")
        .filter(
            fecha_entrada__isnull=True,
            hora_entrada__isnull=True,
            carnet__estado_internado__iexact="FUERA"
        )
        .order_by("-fecha_salida", "-hora_salida")
    )

    no_han_regresado = estudiantes_faltantes.count()


    # =====================================================
    # ÚLTIMOS MOVIMIENTOS
    #
    # Se consultan solamente cuando no existen estudiantes
    # pendientes de regresar.
    # =====================================================

    ultimos_movimientos = []

    if no_han_regresado == 0:

        # Últimas seis salidas del rango
        ultimas_salidas = (
            HistorialInternado.objects
            .select_related("carnet")
            .filter(
                fecha_salida__range=(
                    fecha_desde,
                    fecha_hasta
                )
            )
            .order_by(
                "-fecha_salida",
                "-hora_salida"
            )[:6]
        )

        # Últimas seis entradas del rango
        ultimas_entradas = (
            HistorialInternado.objects
            .select_related("carnet")
            .filter(
                fecha_entrada__range=(
                    fecha_desde,
                    fecha_hasta
                ),
                hora_entrada__isnull=False
            )
            .order_by(
                "-fecha_entrada",
                "-hora_entrada"
            )[:6]
        )


        # Agregar las salidas a una sola lista
        for registro in ultimas_salidas:

            ultimos_movimientos.append({
                "carnet": registro.carnet_id,

                "estudiante": (
                    f"{registro.carnet.nombres} "
                    f"{registro.carnet.apellidos}"
                ),

                "tipo": "Salida",
                "fecha": registro.fecha_salida,
                "hora": registro.hora_salida,
                "estado": "FUERA",

                "fecha_hora": datetime.combine(
                    registro.fecha_salida,
                    registro.hora_salida
                ),
            })


        # Agregar las entradas a la misma lista
        for registro in ultimas_entradas:

            ultimos_movimientos.append({
                "carnet": registro.carnet_id,

                "estudiante": (
                    f"{registro.carnet.nombres} "
                    f"{registro.carnet.apellidos}"
                ),

                "tipo": "Entrada",
                "fecha": registro.fecha_entrada,
                "hora": registro.hora_entrada,
                "estado": "DENTRO",

                "fecha_hora": datetime.combine(
                    registro.fecha_entrada,
                    registro.hora_entrada
                ),
            })


        # Ordenar todos los movimientos por fecha y hora
        ultimos_movimientos.sort(
            key=lambda movimiento: movimiento["fecha_hora"],
            reverse=True
        )

        # Mostrar únicamente los seis más recientes
        ultimos_movimientos = ultimos_movimientos[:6]


    # =====================================================
    # DATOS ENVIADOS AL HTML
    # =====================================================

    context = {
        # Rango seleccionado
        "desde": fecha_desde,
        "hasta": fecha_hasta,

        # Datos generales
        "total_estudiantes": total_estudiantes,
        "estudiantes_dentro": estudiantes_dentro,
        "estudiantes_fuera": estudiantes_fuera,

        # Comedor
        "asistencias_rango": asistencias_rango,
        "desayunos_servidos": desayunos_servidos,
        "almuerzos_servidos": almuerzos_servidos,
        "cenas_servidas": cenas_servidas,

        # Internado
        "salidas_rango": salidas_rango,
        "entradas_rango": entradas_rango,
        "no_han_regresado": no_han_regresado,

        # Tabla inferior
        "estudiantes_faltantes": estudiantes_faltantes,
        "ultimos_movimientos": ultimos_movimientos,
    }

    return render(
        request,
        "dashboard.html",
        context
    )


def testing(request):
  mydata = Estudiantes.objects.all().values()
  template = loader.get_template('prueba.html')
  context = {
    'mymembers': mydata,
  }
  return HttpResponse(template.render(context, request))
 

def index(request):
    error = None # Variable para almacenar el mensaje
    
    if request.method == "POST":
        usuario = request.POST.get("usuario")
        contrasena = request.POST.get("contrasena")
        
        # Consultar si el usuario existe
        user = Usuarios.objects.filter(usuario=usuario, contrasena=contrasena).first()
        
        if user:
            return redirect("dashboard")
        else:
            error = "Usuario o contraseña incorrectos."

    # Si no es POST o si hubo error, volvemos a mostrar el login con el mensaje
    return render(request, "index.html", {"error": error})


def crear_estudiante(request):
    if request.method == "POST":
        foto_url = None

        try:
            carnet = request.POST.get("carnet")
            nombres = request.POST.get("nombres")
            apellidos = request.POST.get("apellidos")
            carrera = request.POST.get("carrera")
            edad = request.POST.get("edad")
            ano = request.POST.get("ano")
            procedencia = request.POST.get("procedencia")
            estado_internado = request.POST.get("estado_internado")

            # 1. Validar si el carnet ya existe ANTES de subir foto
            if Estudiantes.objects.filter(carnet=carnet).exists():
                return JsonResponse({
                    "mensaje": "error",
                    "detalle": "Ya existe un estudiante con ese carnet."
                })

            foto = request.FILES.get("foto")

            # 2. Subir foto solo si el carnet no existe
            if foto:
                foto_url = subir_foto_estudiante(foto, carnet)

            print("Foto URL:", foto_url)

            # 3. Crear estudiante
            estudiante = Estudiantes.objects.create(
                carnet=carnet,
                nombres=nombres,
                apellidos=apellidos,
                carrera=carrera,
                edad_estudi=edad,
                anocarrera=ano,
                procedencia=procedencia,
                foto_estudi=foto_url,
                estado_internado=estado_internado
            )

            print("ESTUDIANTE CREADO:", estudiante.carnet)

            return JsonResponse({
                "mensaje": "ok"
            })

        except Exception as e:
            print("ERROR AL GUARDAR:")
            print(e)

            # Si la foto se subió pero falló guardar en PostgreSQL,
            # intentamos eliminar esa foto para no dejar basura en Storage.
            if foto_url:
                try:
                    eliminar_foto_estudiante(foto_url)
                except Exception as error_foto:
                    print("No se pudo eliminar la foto subida:")
                    print(error_foto)

            return JsonResponse({
                "mensaje": "error",
                "detalle": str(e)
            })

    return JsonResponse({
        "mensaje": "error",
        "detalle": "Método no permitido."
    })
        
def eliminar_estudiante(request):
    if request.method == "POST":
        try:
            carnet = request.POST.get("carnet")

            print("===== ELIMINAR ESTUDIANTE =====")
            print("Carnet recibido:", carnet)

            if not carnet:
                return JsonResponse({
                    "mensaje": "error",
                    "detalle": "No se recibió el carnet."
                })

            estudiante = Estudiantes.objects.get(carnet=carnet)

            foto_url = estudiante.foto_estudi

            if foto_url:
                try:
                 eliminar_foto_estudiante(foto_url)
                except Exception as e:
                    print("No se pudo eliminar la foto del Storage:")
                    print(e)

            estudiante.delete()

            print("ESTUDIANTE ELIMINADO:", carnet)

            return JsonResponse({
                "mensaje": "ok"
            })

        except Estudiantes.DoesNotExist:
            print("ERROR: estudiante no existe")

            return JsonResponse({
                "mensaje": "error",
                "detalle": "El estudiante no existe."
            })

        except IntegrityError as e:
            print("ERROR DE RELACIÓN:")
            print(e)

            return JsonResponse({
                "mensaje": "error",
                "detalle": "No se puede eliminar porque este estudiante tiene datos relacionados, como usuario, asistencia o códigos QR."
            })

        except Exception as e:
            print("ERROR AL ELIMINAR:")
            print(e)

            return JsonResponse({
                "mensaje": "error",
                "detalle": str(e)
            })

    return JsonResponse({
        "mensaje": "error",
        "detalle": "Método no permitido."
    })

def actualizar_estudiante(request):
    if request.method == "POST":
        nueva_foto_url = None

        try:
            carnet = request.POST.get("carnet")
            nombres = request.POST.get("nombres")
            apellidos = request.POST.get("apellidos")
            carrera = request.POST.get("carrera")
            edad = request.POST.get("edad")
            ano = request.POST.get("ano")
            procedencia = request.POST.get("procedencia")
            estado_internado = request.POST.get("estado_internado")
            foto = request.FILES.get("foto")


            estudiante = Estudiantes.objects.get(carnet=carnet)

            foto_anterior_url = estudiante.foto_estudi

            estudiante.nombres = nombres
            estudiante.apellidos = apellidos
            estudiante.carrera = carrera
            estudiante.edad_estudi = edad
            estudiante.anocarrera = ano
            estudiante.procedencia = procedencia
            estudiante.estado_internado = estado_internado

            # Si el usuario seleccionó una nueva foto
            if foto:
                nueva_foto_url = subir_foto_estudiante(foto, carnet)
                estudiante.foto_estudi = nueva_foto_url

            estudiante.save()

            # Si se guardó la nueva foto correctamente, borramos la foto anterior
            if foto and foto_anterior_url:
                try:
                    eliminar_foto_estudiante(foto_anterior_url)
                except Exception as e:
                    print("No se pudo eliminar la foto anterior:")
                    print(e)

            print("ESTUDIANTE ACTUALIZADO:", estudiante.carnet)

            return JsonResponse({
                "mensaje": "ok"
            })

        except Estudiantes.DoesNotExist:
            print("ERROR: estudiante no existe")

            return JsonResponse({
                "mensaje": "error",
                "detalle": "El estudiante no existe."
            })

        except Exception as e:
            print("ERROR AL ACTUALIZAR:")
            print(e)

            # Si se subió una foto nueva pero falló la actualización,
            # borramos esa foto nueva para no dejar basura en Storage.
            if nueva_foto_url:
                try:
                    eliminar_foto_estudiante(nueva_foto_url)
                except Exception as error_foto:
                    print("No se pudo eliminar la nueva foto después del error:")
                    print(error_foto)

            return JsonResponse({
                "mensaje": "error",
                "detalle": str(e)
            })

    return JsonResponse({
        "mensaje": "error",
        "detalle": "Método no permitido."
    })


def estudiantes(request):
    lista_estudiantes = Estudiantes.objects.all().order_by("nombres", "apellidos")

    total_estudiantes = lista_estudiantes.count()
    total_internos = lista_estudiantes.filter(Q(estado_internado__iexact="DENTRO") | Q(estado_internado__iexact="FUERA")).count()
    total_externos = lista_estudiantes.filter(estado_internado__iexact="EXTERNO").count()

    return render(request, "estudiantes.html", {
        "estudiantes": lista_estudiantes,
        "total_estudiantes": total_estudiantes,
        "total_internos": total_internos,
        "total_externos": total_externos
    })


def importar_excel(request):
    if request.method == "POST":
        try:
            archivo = request.FILES.get("archivo")

            if not archivo:
                return JsonResponse({
                    "mensaje": "error",
                    "detalle": "No se recibió ningún archivo."
                })

            if not archivo.name.endswith(".xlsx"):
                return JsonResponse({
                    "mensaje": "error",
                    "detalle": "Solo se permiten archivos Excel .xlsx"
                })

            workbook = load_workbook(archivo)
            hoja = workbook.active

            encabezados = []

            for celda in hoja[1]:
                encabezados.append(str(celda.value).strip().lower())

            columnas_necesarias = [
                "carnet",
                "nombres",
                "apellidos",
                "carrera",
                "edad",
                "ano",
                "procedencia",
                "tipo"
            ]

            for columna in columnas_necesarias:
                if columna not in encabezados:
                    return JsonResponse({
                        "mensaje": "error",
                        "detalle": f"Falta la columna obligatoria: {columna}"
                    })

            creados = 0
            repetidos = 0
            errores = 0

            for fila in hoja.iter_rows(min_row=2, values_only=True):
                try:
                     # Ignorar filas completamente vacías
                    if all(
                        valor is None or str(valor).strip() == ""
                        for valor in fila
                    ):
                        continue
                    datos = dict(zip(encabezados, fila))

                    carnet = str(datos.get("carnet")).strip() if datos.get("carnet") else ""
                    nombres = str(datos.get("nombres")).strip() if datos.get("nombres") else ""
                    apellidos = str(datos.get("apellidos")).strip() if datos.get("apellidos") else ""
                    carrera = str(datos.get("carrera")).strip() if datos.get("carrera") else ""
                    edad = datos.get("edad")
                    ano = str(datos.get("ano")).strip() if datos.get("ano") else ""
                    procedencia = str(datos.get("procedencia")).strip() if datos.get("procedencia") else ""
                    
                    tipo = ( 
                        str(datos.get("tipo")).strip().upper()
                        if datos.get("tipo") else ""
                    )

                    if not carnet or not nombres or not apellidos or not carrera or not edad or not ano:
                        errores += 1
                        continue              

                    # Validar tipo de estudiante
                    if tipo not in ["INTERNO", "EXTERNO", "DENTRO"]:
                        errores += 1
                        continue

                    # Convertir INTERNO en DENTRO
                    if tipo in ["INTERNO", "DENTRO"]:    
                        estado_internado = "DENTRO"
                    else:
                        estado_internado = "EXTERNO"
                    
                    if Estudiantes.objects.filter(carnet=carnet).exists():
                        repetidos += 1
                        continue

                    Estudiantes.objects.create(
                        carnet=carnet,
                        nombres=nombres,
                        apellidos=apellidos,
                        carrera=carrera,
                        edad_estudi=int(edad),
                        anocarrera=ano,
                        procedencia=procedencia,
                        foto_estudi=None,
                        estado_internado=estado_internado
                    )

                    creados += 1

                except Exception as e:
                    print("Error en fila Excel:")
                    print(e)
                    errores += 1

            return JsonResponse({
                "mensaje": "ok",
                "creados": creados,
                "repetidos": repetidos,
                "errores": errores
            })

        except Exception as e:

            return JsonResponse({
                "mensaje": "error",
                "detalle": str(e)
            })

    return JsonResponse({
        "mensaje": "error",
        "detalle": "Método no permitido."
    })




def error(detalle):
    return JsonResponse({"mensaje": "error", "detalle": detalle})


def usuarios(request):
    lista_usuarios = (Usuarios.objects.filter(rol="ESTUDIANTE").select_related("carnet").order_by("id"))
    total_usuarios = lista_usuarios.count()

    return render(request, "usuarios.html", {
        "usuarios": lista_usuarios,
        "total_usuarios": total_usuarios,
    })


def crear_usuario(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        carnet = request.POST.get("carnet", "").strip()
        nombre_usuario = request.POST.get("usuario", "").strip()
        contrasena = request.POST.get("contrasena", "").strip()
        rol = request.POST.get("rol", "").strip()
        estado = request.POST.get("estado") == "true"
        sesion_activa = request.POST.get("sesion_activa") == "true"

        if not nombre_usuario:
            return error("El nombre de usuario es obligatorio.")

        if not contrasena:
            return error("La contraseña es obligatoria.")

        if rol not in ["ESTUDIANTE", "ADMINISTRADOR"]:
            return error("El rol seleccionado no es válido.")

        if Usuarios.objects.filter(usuario=nombre_usuario).exists():
            return error("Ya existe una cuenta con ese nombre de usuario.")

        if GuardaSeguridad.objects.filter(usuario=nombre_usuario).exists():
            return error("Ese nombre de usuario ya pertenece al personal de seguridad.")

        estudiante = None

        if rol == "ESTUDIANTE":
            if not carnet:
                return error("El carnet es obligatorio para una cuenta de estudiante.")

            try:
                estudiante = Estudiantes.objects.get(carnet=carnet)
            except Estudiantes.DoesNotExist:
                return error("No existe ningún estudiante registrado con ese carnet.")

            if Usuarios.objects.filter(carnet_id=carnet).exists():
                return error("Este estudiante ya tiene una cuenta de usuario.")

        Usuarios.objects.create(
            carnet=estudiante,
            usuario=nombre_usuario,
            contrasena=contrasena,
            rol=rol,
            estado=estado,
            sesion_activa=sesion_activa
        )

        return JsonResponse({"mensaje": "ok"})

    except Exception as e:
        return error(str(e))


def actualizar_usuario(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        usuario_id = request.POST.get("id", "").strip()
        carnet = request.POST.get("carnet", "").strip()
        nombre_usuario = request.POST.get("usuario", "").strip()
        rol = request.POST.get("rol", "").strip()
        estado = request.POST.get("estado") == "true"
        sesion_activa = request.POST.get("sesion_activa") == "true"

        if not usuario_id:
            return error("No se recibió el ID del usuario.")

        if not nombre_usuario:
            return error("El nombre de usuario es obligatorio.")

        if rol not in ["ESTUDIANTE", "ADMINISTRADOR"]:
            return error("El rol seleccionado no es válido.")

        usuario = Usuarios.objects.get(id=usuario_id)

        if usuario.rol in ["ADMINISTRADOR", "ADMIN"]:
            return error("No se puede editar una cuenta de administrador desde esta sección.")

        if Usuarios.objects.filter(usuario=nombre_usuario).exclude(id=usuario_id).exists():
            return error("Ya existe otra cuenta con ese nombre de usuario.")

        if GuardaSeguridad.objects.filter(usuario=nombre_usuario).exists():
            return error("Ese nombre de usuario ya pertenece al personal de seguridad.")

        estudiante = None

        if rol == "ESTUDIANTE":
            if not carnet:
                return error("El carnet es obligatorio para una cuenta de estudiante.")

            try:
                estudiante = Estudiantes.objects.get(carnet=carnet)
            except Estudiantes.DoesNotExist:
                return error("No existe ningún estudiante registrado con ese carnet.")

            if Usuarios.objects.filter(carnet_id=carnet).exclude(id=usuario_id).exists():
                return error("Ese estudiante ya está vinculado a otra cuenta.")

        usuario.carnet = estudiante
        usuario.usuario = nombre_usuario
        usuario.rol = rol
        usuario.estado = estado
        usuario.sesion_activa = sesion_activa
        usuario.save()

        return JsonResponse({"mensaje": "ok"})

    except Usuarios.DoesNotExist:
        return error("El usuario no existe.")

    except Exception as e:
        return error(str(e))


def actualizar_password_usuario(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        usuario_id = request.POST.get("id", "").strip()
        contrasena = request.POST.get("contrasena", "").strip()

        if not usuario_id:
            return error("No se recibió el ID del usuario.")

        if not contrasena:
            return error("La nueva contraseña es obligatoria.")

        usuario = Usuarios.objects.get(id=usuario_id)

        if usuario.rol in ["ADMINISTRADOR", "ADMIN"]:
            return error("No se puede cambiar la contraseña de un administrador desde esta sección.")

        usuario.contrasena = contrasena
        usuario.save()

        return JsonResponse({"mensaje": "ok"})

    except Usuarios.DoesNotExist:
        return error("El usuario no existe.")

    except Exception as e:
        return error(str(e))


def eliminar_usuario(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        usuario_id = request.POST.get("id", "").strip()

        if not usuario_id:
            return error("No se recibió el ID del usuario.")

        usuario = Usuarios.objects.get(id=usuario_id)

        if usuario.rol in ["ADMINISTRADOR", "ADMIN"]:
            return error("No se puede eliminar una cuenta de administrador.")

        usuario.delete()

        return JsonResponse({"mensaje": "ok"})

    except Usuarios.DoesNotExist:
        return error("El usuario no existe.")

    except Exception as e:
        return error(str(e))





def guardas(request):
    lista_guardas = GuardaSeguridad.objects.all().order_by("nombre")

    return render(request, "guardas.html", {
        "guardas_seguridad": lista_guardas,
        "total_seguridad": lista_guardas.count()
    })


def crear_guarda_seguridad(request):
    if request.method != "POST":
        return error("Método no permitido.")

    foto_url = None

    try:
        codigo = request.POST.get("codigo_unico", "").strip()
        nombre = request.POST.get("nombre", "").strip()
        apellido = request.POST.get("apellido", "").strip()
        sexo = request.POST.get("sexo", "").strip()
        turno = request.POST.get("turno", "").strip()
        foto = request.FILES.get("foto")

        if not codigo:
            return error("El código único es obligatorio.")

        if not nombre:
            return error("El nombre del guarda es obligatorio.")

        if not apellido:
            return error("El apellido del guarda es obligatorio.")

        if not sexo:
            return error("Debe seleccionar el sexo.")

        if not turno:
            return error("Debe seleccionar el turno.")

        if GuardaSeguridad.objects.filter(codigo_unico=codigo).exists():
            return error("Ya existe un guarda registrado con ese código.")

        primer_nombre = nombre.split()[0]
        primer_apellido = apellido.split()[0]

        usuario = f"{primer_nombre}.{primer_apellido}".lower()
        usuario = unicodedata.normalize("NFKD", usuario).encode("ascii", "ignore").decode("ascii")
        usuario = "".join(c for c in usuario if c.isalnum() or c == ".")

        usuario_base = usuario
        numero = 1

        while GuardaSeguridad.objects.filter(usuario=usuario).exists() or Usuarios.objects.filter(usuario=usuario).exists():
            usuario = f"{usuario_base}{numero}"
            numero += 1

        caracteres = string.ascii_letters + string.digits
        contrasena = "".join(secrets.choice(caracteres) for _ in range(10))

        if foto:
            foto_url = subir_foto_guarda(foto, codigo)

        GuardaSeguridad.objects.create(
            codigo_unico=codigo,
            nombre=nombre,
            apellido=apellido,
            sexo=sexo,
            foto=foto_url,
            turno=turno,
            usuario=usuario,
            contrasena=contrasena
        )

        return JsonResponse({"mensaje": "ok", "usuario": usuario, "contrasena": contrasena})

    except Exception as e:
        if foto_url:
            try:
                eliminar_foto_guarda(foto_url)
            except Exception as error_foto:
                print("No se pudo eliminar la foto subida:", error_foto)

        return error(str(e))


def actualizar_guarda_seguridad(request):
    if request.method != "POST":
        return error("Método no permitido.")

    nueva_foto_url = None

    try:
        codigo = request.POST.get("codigo_unico", "").strip()
        nombre = request.POST.get("nombre", "").strip()
        apellido = request.POST.get("apellido", "").strip()
        sexo = request.POST.get("sexo", "").strip()
        turno = request.POST.get("turno", "").strip()
        foto = request.FILES.get("foto")

        if not codigo:
            return error("No se recibió el código del guarda.")

        if not nombre:
            return error("El nombre del guarda es obligatorio.")

        if not apellido:
            return error("El apellido del guarda es obligatorio.")

        if not sexo:
            return error("Debe seleccionar el sexo.")

        if not turno:
            return error("Debe seleccionar el turno.")

        guarda = GuardaSeguridad.objects.get(codigo_unico=codigo)
        foto_anterior_url = guarda.foto

        guarda.nombre = nombre
        guarda.apellido = apellido
        guarda.sexo = sexo
        guarda.turno = turno

        if foto:
            nueva_foto_url = subir_foto_guarda(foto, codigo)
            guarda.foto = nueva_foto_url

        guarda.save()

        if foto and foto_anterior_url:
            try:
                eliminar_foto_guarda(foto_anterior_url)
            except Exception as e:
                print("No se pudo eliminar la foto anterior:", e)

        return JsonResponse({"mensaje": "ok"})

    except GuardaSeguridad.DoesNotExist:
        return error("El guarda de seguridad no existe.")

    except Exception as e:
        if nueva_foto_url:
            try:
                eliminar_foto_guarda(nueva_foto_url)
            except Exception as error_foto:
                print("No se pudo eliminar la nueva foto:", error_foto)

        return error(str(e))
    

def actualizar_password_guarda(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        codigo = request.POST.get("codigo_unico", "").strip()
        contrasena = request.POST.get("contrasena", "").strip()

        if not codigo:
            return error("No se recibió el código del guarda.")

        if not contrasena:
            return error("La nueva contraseña es obligatoria.")

        guarda = GuardaSeguridad.objects.get(codigo_unico=codigo)
        guarda.contrasena = contrasena
        guarda.save()

        return JsonResponse({"mensaje": "ok"})

    except GuardaSeguridad.DoesNotExist:
        return error("El guarda de seguridad no existe.")

    except Exception as e:
        return error(str(e))


def eliminar_guarda_seguridad(request):
    if request.method != "POST":
        return error("Método no permitido.")

    try:
        codigo = request.POST.get("codigo_unico", "").strip()

        if not codigo:
            return error("No se recibió el código del guarda.")

        guarda = GuardaSeguridad.objects.get(codigo_unico=codigo)
        foto_url = guarda.foto

        guarda.delete()

        if foto_url:
            try:
                eliminar_foto_guarda(foto_url)
            except Exception as e:
                print("No se pudo eliminar la foto del Storage:", e)

        return JsonResponse({"mensaje": "ok"})

    except GuardaSeguridad.DoesNotExist:
        return error("El guarda de seguridad no existe.")

    except Exception as e:
        return error(str(e))






def filtrar_asistencias_reportes(request):
    hoy = timezone.localdate()

    desde = request.GET.get("desde") or hoy.isoformat()
    hasta = request.GET.get("hasta") or hoy.isoformat()
    tipo_comida = request.GET.get("tipo_comida", "").strip()
    carrera = request.GET.get("carrera", "").strip()
    ano = request.GET.get("ano", "").strip()

    try:
        fecha_desde = date.fromisoformat(desde)
        fecha_hasta = date.fromisoformat(hasta)
    except ValueError:
        fecha_desde = hoy
        fecha_hasta = hoy
        desde = hoy.isoformat()
        hasta = hoy.isoformat()

    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde
        desde = fecha_desde.isoformat()
        hasta = fecha_hasta.isoformat()

    asistencias = (
        Asistencias.objects
        .select_related("carnet")
        .filter(fecha__range=[fecha_desde, fecha_hasta])
        .order_by("-fecha", "-hora")
    )

    if tipo_comida:
        asistencias = asistencias.filter(tipo_comida__iexact=tipo_comida)

    if carrera:
        asistencias = asistencias.filter(carnet__carrera=carrera)

    if ano:
        asistencias = asistencias.filter(carnet__anocarrera=ano)

    return asistencias, desde, hasta, tipo_comida, carrera, ano






def obtener_datos_reporte_comedor(request):
    hoy = timezone.localdate()
    primer_dia = hoy.replace(day=1)
    ultimo_dia = hoy.replace(day=calendar.monthrange(hoy.year, hoy.month)[1])

    desde = request.GET.get("desde") or primer_dia.isoformat()
    hasta = request.GET.get("hasta") or ultimo_dia.isoformat()
    carrera = request.GET.get("carrera", "").strip()
    ano = request.GET.get("ano", "").strip()

    try:
        fecha_desde = date.fromisoformat(desde)
        fecha_hasta = date.fromisoformat(hasta)
    except ValueError:
        fecha_desde = primer_dia
        fecha_hasta = ultimo_dia
        desde = fecha_desde.isoformat()
        hasta = fecha_hasta.isoformat()

    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde
        desde = fecha_desde.isoformat()
        hasta = fecha_hasta.isoformat()

    estudiantes = Estudiantes.objects.all().order_by("nombres", "apellidos")

    if carrera:
        estudiantes = estudiantes.filter(carrera=carrera)

    if ano:
        estudiantes = estudiantes.filter(anocarrera=ano)

    estudiantes = list(estudiantes)
    carnets = [e.carnet for e in estudiantes]

    dias_periodo = (fecha_hasta - fecha_desde).days + 1

    conteos = (
        Asistencias.objects
        .filter(carnet_id__in=carnets, fecha__range=[fecha_desde, fecha_hasta], estado=True)
        .values("carnet_id", "tipo_comida")
        .annotate(total=Count("id"))
    )

    mapa = {
        carnet: {
            "DESAYUNO": 0,
            "ALMUERZO": 0,
            "CENA": 0
        }
        for carnet in carnets
    }

    for fila in conteos:
        carnet = fila["carnet_id"]
        tipo = (fila["tipo_comida"] or "").strip().upper()

        if carnet in mapa and tipo in mapa[carnet]:
            mapa[carnet][tipo] += fila["total"]

    reporte_estudiantes = []

    suma_desayunos = 0
    suma_almuerzos = 0
    suma_cenas = 0
    estudiantes_bajo_60 = 0

    for estudiante in estudiantes:
        datos = mapa.get(estudiante.carnet, {})

        desayunos = datos.get("DESAYUNO", 0)
        almuerzos = datos.get("ALMUERZO", 0)
        cenas = datos.get("CENA", 0)

        total = desayunos + almuerzos + cenas

        porcentaje_desayuno = round((desayunos / dias_periodo) * 100, 1)
        porcentaje_almuerzo = round((almuerzos / dias_periodo) * 100, 1)
        porcentaje_cena = round((cenas / dias_periodo) * 100, 1)
        porcentaje_total = round((total / (dias_periodo * 3)) * 100, 1)

        if porcentaje_total < 60:
            estado = "Bajo uso"
            estudiantes_bajo_60 += 1
        elif porcentaje_total >= 85:
            estado = "Excelente"
        else:
            estado = "Normal"

        reporte_estudiantes.append({
            "carnet": estudiante.carnet,
            "estudiante": f"{estudiante.nombres} {estudiante.apellidos}",
            "carrera": estudiante.carrera,
            "ano": estudiante.anocarrera,
            "desayunos": desayunos,
            "almuerzos": almuerzos,
            "cenas": cenas,
            "total": total,
            "porcentaje_desayuno": porcentaje_desayuno,
            "porcentaje_almuerzo": porcentaje_almuerzo,
            "porcentaje_cena": porcentaje_cena,
            "porcentaje_total": porcentaje_total,
            "estado": estado,
        })

        suma_desayunos += desayunos
        suma_almuerzos += almuerzos
        suma_cenas += cenas

    total_estudiantes = len(estudiantes)
    posibles_por_tiempo = total_estudiantes * dias_periodo
    posibles_totales = posibles_por_tiempo * 3

    promedio_desayuno = round((suma_desayunos / posibles_por_tiempo) * 100, 1) if posibles_por_tiempo else 0
    promedio_almuerzo = round((suma_almuerzos / posibles_por_tiempo) * 100, 1) if posibles_por_tiempo else 0
    promedio_cena = round((suma_cenas / posibles_por_tiempo) * 100, 1) if posibles_por_tiempo else 0

    total_asistencias = suma_desayunos + suma_almuerzos + suma_cenas
    promedio_mensual = round((total_asistencias / posibles_totales) * 100, 1) if posibles_totales else 0

    tiempos = {
        "Desayuno": promedio_desayuno,
        "Almuerzo": promedio_almuerzo,
        "Cena": promedio_cena,
    }

    tiempo_menor_uso = min(tiempos, key=tiempos.get) if total_estudiantes else "Sin datos"

    carreras = (
        Estudiantes.objects
        .exclude(carrera__isnull=True)
        .exclude(carrera__exact="")
        .values_list("carrera", flat=True)
        .distinct()
        .order_by("carrera")
    )

    return {
        "reporte_estudiantes": reporte_estudiantes,
        "desde": desde,
        "hasta": hasta,
        "mes": fecha_desde.strftime("%Y-%m"),
        "carrera": carrera,
        "ano": ano,
        "carreras": carreras,
        "dias_periodo": dias_periodo,
        "total_estudiantes": total_estudiantes,
        "estudiantes_bajo_60": estudiantes_bajo_60,
        "promedio_mensual": promedio_mensual,
        "promedio_desayuno": promedio_desayuno,
        "promedio_almuerzo": promedio_almuerzo,
        "promedio_cena": promedio_cena,
        "tiempo_menor_uso": tiempo_menor_uso,
        "total_asistencias": total_asistencias,
    }


def reportes(request):
    datos = obtener_datos_reporte_comedor(request)

    paginador = Paginator(datos["reporte_estudiantes"], 8)
    pagina = request.GET.get("page")
    estudiantes_pagina = paginador.get_page(pagina)

    datos["estudiantes_pagina"] = estudiantes_pagina

    return render(request, "reportes.html", datos)


def exportar_reporte_excel(request):
    datos = obtener_datos_reporte_comedor(request)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Uso del Comedor"

    hoja.append([
        "Carnet",
        "Estudiante",
        "Carrera",
        "Año",
        "Desayunos",
        "% Desayuno",
        "Almuerzos",
        "% Almuerzo",
        "Cenas",
        "% Cena",
        "Total Asistencias",
        "% Total",
        "Estado"
    ])

    for estudiante in datos["reporte_estudiantes"]:
        hoja.append([
            estudiante["carnet"],
            estudiante["estudiante"],
            estudiante["carrera"],
            estudiante["ano"],
            estudiante["desayunos"],
            estudiante["porcentaje_desayuno"],
            estudiante["almuerzos"],
            estudiante["porcentaje_almuerzo"],
            estudiante["cenas"],
            estudiante["porcentaje_cena"],
            estudiante["total"],
            estudiante["porcentaje_total"],
            estudiante["estado"]
        ])

    for columna in hoja.columns:
        max_length = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[columna[0].column_letter].width = max_length + 3

    nombre_archivo = f"uso_comedor_{datos['desde']}_a_{datos['hasta']}.xlsx"

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'

    workbook.save(response)

    return response



def dashboard_comedor(request):
    
    hoy = timezone.localdate()
    inicio_mes = hoy.replace(day=1)
    ayer = hoy - timedelta(days=1)
    inicio_3_dias = ayer - timedelta(days=2)

    estudiantes = Estudiantes.objects.all()
    carnets = list(estudiantes.values_list("carnet", flat=True))
    total_estudiantes = len(carnets)

    # =========================================================
    # ASISTENCIAS DE HOY
    # =========================================================

    asistencias_hoy = Asistencias.objects.filter(fecha=hoy, estado=True)

    resumen_hoy = asistencias_hoy.aggregate(
        desayunos=Count("id", filter=Q(tipo_comida__iexact="DESAYUNO")),
        almuerzos=Count("id", filter=Q(tipo_comida__iexact="ALMUERZO")),
        cenas=Count("id", filter=Q(tipo_comida__iexact="CENA"))
    )

    total_desayunos = resumen_hoy["desayunos"] or 0
    total_almuerzos = resumen_hoy["almuerzos"] or 0
    total_cenas = resumen_hoy["cenas"] or 0
    total_hoy = total_desayunos + total_almuerzos + total_cenas

    porcentaje_desayuno = round((total_desayunos / total_estudiantes) * 100) if total_estudiantes else 0
    porcentaje_almuerzo = round((total_almuerzos / total_estudiantes) * 100) if total_estudiantes else 0
    porcentaje_cena = round((total_cenas / total_estudiantes) * 100) if total_estudiantes else 0

    total_posibles_hoy = total_estudiantes * 3
    porcentaje_total_hoy = round((total_hoy / total_posibles_hoy) * 100) if total_posibles_hoy else 0

    # =========================================================
    # ESTUDIANTES SIN ASISTENCIA HOY
    # =========================================================

    carnets_hoy = set(asistencias_hoy.values_list("carnet_id", flat=True).distinct())
    estudiantes_sin_asistencia_hoy = total_estudiantes - len(carnets_hoy)

    # =========================================================
    # RESUMEN DEL MES ACTUAL
    # =========================================================

    asistencias_mes = Asistencias.objects.filter(fecha__range=[inicio_mes, hoy], estado=True)
    total_asistencias_mes = asistencias_mes.count()
    dias_transcurridos = hoy.day

    posibles_mes = total_estudiantes * dias_transcurridos * 3

    porcentaje_promedio_mes = round((total_asistencias_mes / posibles_mes) * 100, 1) if posibles_mes else 0
    promedio_diario = round(total_asistencias_mes / dias_transcurridos, 1) if dias_transcurridos else 0

    # =========================================================
    # ESTUDIANTES BAJO 60% ESTE MES
    # =========================================================

    uso_por_estudiante = dict(
        asistencias_mes.values("carnet_id")
        .annotate(total=Count("id"))
        .values_list("carnet_id", "total")
    )

    comidas_posibles_estudiante = dias_transcurridos * 3
    estudiantes_bajo_60 = 0

    for carnet in carnets:
        utilizadas = uso_por_estudiante.get(carnet, 0)
        porcentaje = (utilizadas / comidas_posibles_estudiante) * 100 if comidas_posibles_estudiante else 0

        if porcentaje < 60:
            estudiantes_bajo_60 += 1

    # =========================================================
    # 3 DÍAS CONSECUTIVOS SIN NINGUNA ASISTENCIA
    # Se usan los últimos 3 días COMPLETOS, sin contar hoy
    # =========================================================

    registros_3_dias = set(
        Asistencias.objects
        .filter(fecha__range=[inicio_3_dias, ayer], estado=True)
        .values_list("carnet_id", "fecha")
        .distinct()
    )

    estudiantes_3_dias_ausentes = 0

    for carnet in carnets:
        ausente_los_3 = True

        for i in range(3):
            fecha = inicio_3_dias + timedelta(days=i)

            if (carnet, fecha) in registros_3_dias:
                ausente_los_3 = False
                break

        if ausente_los_3:
            estudiantes_3_dias_ausentes += 1

    # =========================================================
    # ASISTENCIAS DE HOY POR CARRERA
    # =========================================================

    estudiantes_por_carrera = (
        Estudiantes.objects
        .exclude(carrera__isnull=True)
        .exclude(carrera__exact="")
        .values("carrera")
        .annotate(estudiantes=Count("carnet"))
        .order_by("carrera")
    )

    asistencias_por_carrera = {
        item["carnet__carrera"]: item["total"]
        for item in asistencias_hoy
        .values("carnet__carrera")
        .annotate(total=Count("id"))
    }

    carreras_resumen = []

    for item in estudiantes_por_carrera:
        nombre_carrera = item["carrera"]
        cantidad_estudiantes = item["estudiantes"]
        posibles = cantidad_estudiantes * 3
        asistencias = asistencias_por_carrera.get(nombre_carrera, 0)
        porcentaje = round((asistencias / posibles) * 100) if posibles else 0

        carreras_resumen.append({
            "carrera": nombre_carrera,
            "asistencias": asistencias,
            "posibles": posibles,
            "porcentaje": porcentaje
        })

    # =========================================================
    # ÚLTIMAS ASISTENCIAS
    # =========================================================

    ultimas_asistencias = (
        Asistencias.objects
        .select_related("carnet")
        .filter(estado=True)
        .order_by("-fecha", "-hora")[:6]
    )

    context = {
        "hoy": hoy,
        "inicio_mes": inicio_mes,

        "total_estudiantes": total_estudiantes,

        "total_desayunos": total_desayunos,
        "total_almuerzos": total_almuerzos,
        "total_cenas": total_cenas,
        "total_hoy": total_hoy,

        "porcentaje_desayuno": porcentaje_desayuno,
        "porcentaje_almuerzo": porcentaje_almuerzo,
        "porcentaje_cena": porcentaje_cena,
        "porcentaje_total_hoy": porcentaje_total_hoy,

        "total_posibles_hoy": total_posibles_hoy,

        "estudiantes_sin_asistencia_hoy": estudiantes_sin_asistencia_hoy,
        "estudiantes_bajo_60": estudiantes_bajo_60,
        "estudiantes_3_dias_ausentes": estudiantes_3_dias_ausentes,

        "dias_transcurridos": dias_transcurridos,
        "total_asistencias_mes": total_asistencias_mes,
        "promedio_diario": promedio_diario,
        "porcentaje_promedio_mes": porcentaje_promedio_mes,

        "carreras_resumen": carreras_resumen[:5],
        "ultimas_asistencias": ultimas_asistencias,
    }

    return render(request, "dashboard_comedor.html", context)


def formatear_tiempo_fuera(inicio, fin):
    if not inicio or not fin:
        return "—"

    segundos = max(int((fin - inicio).total_seconds()), 0)
    dias, resto = divmod(segundos, 86400)
    horas, resto = divmod(resto, 3600)
    minutos = resto // 60

    if dias > 0:
        return f"{dias} d {horas} h {minutos} min"

    if horas > 0:
        return f"{horas} h {minutos} min"

    return f"{minutos} min"


def obtener_datos_reporte_internado(request):
    hoy = timezone.localdate()
    primer_dia = hoy.replace(day=1)

    desde = request.GET.get("desde") or primer_dia.isoformat()
    hasta = request.GET.get("hasta") or hoy.isoformat()
    carrera = request.GET.get("carrera", "").strip()
    ano = request.GET.get("ano", "").strip()
    estado = request.GET.get("estado", "").strip().upper()

    try:
        fecha_desde = date.fromisoformat(desde)
        fecha_hasta = date.fromisoformat(hasta)
    except ValueError:
        fecha_desde = primer_dia
        fecha_hasta = hoy
        desde = fecha_desde.isoformat()
        hasta = fecha_hasta.isoformat()

    if fecha_desde > fecha_hasta:
        fecha_desde, fecha_hasta = fecha_hasta, fecha_desde
        desde = fecha_desde.isoformat()
        hasta = fecha_hasta.isoformat()

    estudiantes = Estudiantes.objects.filter(Q(estado_internado__iexact="DENTRO") | Q(estado_internado__iexact="FUERA"))

    if carrera:
        estudiantes = estudiantes.filter(carrera=carrera)

    if ano:
        estudiantes = estudiantes.filter(anocarrera=ano)

    if estado:
        estudiantes = estudiantes.filter(estado_internado__iexact=estado)

    estudiantes = list(estudiantes.order_by("nombres", "apellidos"))
    carnets = [e.carnet for e in estudiantes]

    movimientos = list(
        HistorialInternado.objects
        .filter(carnet_id__in=carnets)
        .filter(Q(fecha_salida__range=[fecha_desde, fecha_hasta]) | Q(fecha_entrada__range=[fecha_desde, fecha_hasta]))
        .order_by("fecha_salida", "hora_salida")
    )

    pendientes_qs = list(
        HistorialInternado.objects
        .select_related("carnet")
        .filter(carnet_id__in=carnets, fecha_salida__lte=fecha_hasta)
        .filter(Q(fecha_entrada__isnull=True) | Q(fecha_entrada__gt=fecha_hasta))
        .order_by("fecha_salida", "hora_salida")
    )

    mapa_pendientes = {}

    for movimiento in pendientes_qs:
        mapa_pendientes[movimiento.carnet_id] = movimiento

    mapa = {}

    for estudiante in estudiantes:
        mapa[estudiante.carnet] = {
            "estudiante": f"{estudiante.nombres} {estudiante.apellidos}",
            "carnet": estudiante.carnet,
            "carrera": estudiante.carrera,
            "ano": estudiante.anocarrera,
            "estado": estudiante.estado_internado.upper(),
            "salidas": 0,
            "regresos": 0,
            "ultima_salida_fecha": None,
            "ultima_salida_hora": None,
            "ultimo_regreso_fecha": None,
            "ultimo_regreso_hora": None,
        }

    total_salidas = 0
    total_regresos = 0

    for movimiento in movimientos:
        datos = mapa.get(movimiento.carnet_id)

        if not datos:
            continue

        if fecha_desde <= movimiento.fecha_salida <= fecha_hasta:
            datos["salidas"] += 1
            total_salidas += 1

            if not datos["ultima_salida_fecha"] or (movimiento.fecha_salida, movimiento.hora_salida) > (datos["ultima_salida_fecha"], datos["ultima_salida_hora"]):
                datos["ultima_salida_fecha"] = movimiento.fecha_salida
                datos["ultima_salida_hora"] = movimiento.hora_salida

        if movimiento.fecha_entrada and fecha_desde <= movimiento.fecha_entrada <= fecha_hasta:
            datos["regresos"] += 1
            total_regresos += 1

            if not datos["ultimo_regreso_fecha"] or (movimiento.fecha_entrada, movimiento.hora_entrada) > (datos["ultimo_regreso_fecha"], datos["ultimo_regreso_hora"]):
                datos["ultimo_regreso_fecha"] = movimiento.fecha_entrada
                datos["ultimo_regreso_hora"] = movimiento.hora_entrada

    zona = timezone.get_current_timezone()

    if fecha_hasta >= hoy:
        fecha_referencia = timezone.localtime()
    else:
        fecha_referencia = timezone.make_aware(datetime.combine(fecha_hasta, time(23, 59, 59)), zona)

    resumen_estudiantes = []
    pendientes = []

    for datos in mapa.values():
        pendiente = mapa_pendientes.get(datos["carnet"])

        if datos["salidas"] == 0 and datos["regresos"] == 0 and not pendiente:
            continue

        ultima_salida = "—"
        ultimo_regreso = "—"
        tiempo_fuera = "—"

        if datos["ultima_salida_fecha"]:
            ultima_salida = f'{datos["ultima_salida_fecha"].strftime("%d/%m/%Y")} {datos["ultima_salida_hora"].strftime("%H:%M")}'

        if datos["ultimo_regreso_fecha"]:
            ultimo_regreso = f'{datos["ultimo_regreso_fecha"].strftime("%d/%m/%Y")} {datos["ultimo_regreso_hora"].strftime("%H:%M")}'

        if pendiente:
            salida_dt = timezone.make_aware(datetime.combine(pendiente.fecha_salida, pendiente.hora_salida), zona)
            tiempo_fuera = formatear_tiempo_fuera(salida_dt, fecha_referencia)
            ultima_salida = f'{pendiente.fecha_salida.strftime("%d/%m/%Y")} {pendiente.hora_salida.strftime("%H:%M")}'

            pendientes.append({
                "estudiante": datos["estudiante"],
                "carnet": datos["carnet"],
                "carrera": datos["carrera"],
                "ultima_salida": ultima_salida,
                "tiempo_fuera": tiempo_fuera,
            })

        resumen_estudiantes.append({
            "estudiante": datos["estudiante"],
            "carnet": datos["carnet"],
            "carrera": datos["carrera"],
            "ano": datos["ano"],
            "estado": datos["estado"],
            "salidas": datos["salidas"],
            "regresos": datos["regresos"],
            "ultima_salida": ultima_salida,
            "ultimo_regreso": ultimo_regreso,
            "pendiente": pendiente is not None,
            "tiempo_fuera": tiempo_fuera,
        })

    carreras = (
        Estudiantes.objects
        .exclude(carrera__isnull=True)
        .exclude(carrera="")
        .values_list("carrera", flat=True)
        .distinct()
        .order_by("carrera")
    )

    return {
        "resumen_estudiantes": resumen_estudiantes,
        "pendientes": pendientes,
        "desde": desde,
        "hasta": hasta,
        "carrera": carrera,
        "ano": ano,
        "estado": estado,
        "carreras": carreras,
        "total_salidas": total_salidas,
        "total_regresos": total_regresos,
        "total_movimientos": total_salidas + total_regresos,
        "total_pendientes": len(pendientes),
    }


def reporte_internado(request):
    datos = obtener_datos_reporte_internado(request)
    paginador = Paginator(datos["resumen_estudiantes"], 8)
    datos["estudiantes_pagina"] = paginador.get_page(request.GET.get("page"))
    return render(request, "reporte_internado.html", datos)


def exportar_reporte_internado_excel(request):
    datos = obtener_datos_reporte_internado(request)

    workbook = Workbook()
    hoja = workbook.active
    hoja.title = "Movimientos Internado"

    hoja.append([
        "Carnet",
        "Estudiante",
        "Carrera",
        "Año",
        "Salidas",
        "Regresos",
        "Última salida",
        "Último regreso",
        "Tiempo fuera",
        "Estado"
    ])

    for r in datos["resumen_estudiantes"]:
        hoja.append([
            r["carnet"],
            r["estudiante"],
            r["carrera"],
            r["ano"],
            r["salidas"],
            r["regresos"],
            r["ultima_salida"],
            r["ultimo_regreso"],
            r["tiempo_fuera"],
            "NO HA REGRESADO" if r["pendiente"] else r["estado"]
        ])

    for columna in hoja.columns:
        max_length = max(len(str(celda.value or "")) for celda in columna)
        hoja.column_dimensions[columna[0].column_letter].width = max_length + 3

    nombre = f"reporte_internado_{datos['desde']}_a_{datos['hasta']}.xlsx"

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nombre}"'

    workbook.save(response)

    return response




def dashboard_internado(request):
    hoy = timezone.localdate()
    ahora = timezone.localtime()
    zona = timezone.get_current_timezone()

    internos = Estudiantes.objects.filter(Q(estado_internado__iexact="DENTRO") | Q(estado_internado__iexact="FUERA"))

    total_internos = internos.count()
    total_dentro = internos.filter(estado_internado__iexact="DENTRO").count()
    total_fuera = internos.filter(estado_internado__iexact="FUERA").count()

    salidas_hoy = HistorialInternado.objects.filter(fecha_salida=hoy).count()
    regresos_hoy = HistorialInternado.objects.filter(fecha_entrada=hoy).count()

    carnets_fuera = list(
        internos
        .filter(estado_internado__iexact="FUERA")
        .values_list("carnet", flat=True)
    )

    movimientos_pendientes = (
        HistorialInternado.objects
        .select_related("carnet")
        .filter(carnet_id__in=carnets_fuera, fecha_entrada__isnull=True)
        .order_by("-fecha_salida", "-hora_salida")
    )

    pendientes_por_carnet = {}

    for movimiento in movimientos_pendientes:
        if movimiento.carnet_id not in pendientes_por_carnet:
            pendientes_por_carnet[movimiento.carnet_id] = movimiento

    pendientes = []
    total_mas_24_horas = 0
    total_mas_3_dias = 0

    for movimiento in pendientes_por_carnet.values():
        salida = timezone.make_aware(datetime.combine(movimiento.fecha_salida, movimiento.hora_salida), zona)
        diferencia = ahora - salida
        segundos = max(int(diferencia.total_seconds()), 0)

        dias, resto = divmod(segundos, 86400)
        horas, resto = divmod(resto, 3600)
        minutos = resto // 60

        if dias:
            tiempo_fuera = f"{dias} d {horas} h"
        elif horas:
            tiempo_fuera = f"{horas} h {minutos} min"
        else:
            tiempo_fuera = f"{minutos} min"

        if diferencia >= timedelta(days=3):
            nivel = "critico"
            total_mas_3_dias += 1
            total_mas_24_horas += 1
        elif diferencia >= timedelta(hours=24):
            nivel = "alerta"
            total_mas_24_horas += 1
        else:
            nivel = "normal"

        pendientes.append({
            "estudiante": f"{movimiento.carnet.nombres} {movimiento.carnet.apellidos}",
            "carnet": movimiento.carnet.carnet,
            "carrera": movimiento.carnet.carrera,
            "salida": f"{movimiento.fecha_salida.strftime('%d/%m/%Y')} {movimiento.hora_salida.strftime('%H:%M')}",
            "tiempo_fuera": tiempo_fuera,
            "nivel": nivel,
        })

    pendientes.sort(key=lambda x: 0 if x["nivel"] == "critico" else 1 if x["nivel"] == "alerta" else 2)

    historiales = (
        HistorialInternado.objects
        .select_related("carnet")
        .order_by("-fecha_salida", "-hora_salida")[:8]
    )

    movimientos = []

    for h in historiales:
        movimientos.append({
            "estudiante": f"{h.carnet.nombres} {h.carnet.apellidos}",
            "carnet": h.carnet.carnet,
            "tipo": "SALIDA",
            "fecha": h.fecha_salida.strftime("%d/%m/%Y"),
            "hora": h.hora_salida.strftime("%H:%M"),
            "orden": datetime.combine(h.fecha_salida, h.hora_salida),
        })

        if h.fecha_entrada and h.hora_entrada:
            movimientos.append({
                "estudiante": f"{h.carnet.nombres} {h.carnet.apellidos}",
                "carnet": h.carnet.carnet,
                "tipo": "ENTRADA",
                "fecha": h.fecha_entrada.strftime("%d/%m/%Y"),
                "hora": h.hora_entrada.strftime("%H:%M"),
                "orden": datetime.combine(h.fecha_entrada, h.hora_entrada),
            })

    movimientos.sort(key=lambda x: x["orden"], reverse=True)
    ultimos_movimientos = movimientos[:8]

    return render(request, "dashboard_internado.html", {
        "hoy": hoy,
        "total_internos": total_internos,
        "total_dentro": total_dentro,
        "total_fuera": total_fuera,
        "salidas_hoy": salidas_hoy,
        "regresos_hoy": regresos_hoy,
        "total_mas_24_horas": total_mas_24_horas,
        "total_mas_3_dias": total_mas_3_dias,
        "pendientes": pendientes,
        "ultimos_movimientos": ultimos_movimientos,
    })


